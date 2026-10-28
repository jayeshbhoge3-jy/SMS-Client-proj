from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib import messages
from django.contrib.auth import update_session_auth_hash
from company.models import Transaction
from accounts.models import UserProfile
from django.core.paginator import Paginator
from django.shortcuts import  get_object_or_404
from django.http import HttpResponse
from django.core.mail import EmailMultiAlternatives
from django.template.loader import render_to_string
from reportlab.lib.pagesizes import A4
from datetime import date, timedelta
from reportlab.pdfgen import canvas
from company.models import Company, Staff
from inventory.models import StockTransaction, Product, ProductBatch, SalesOrder 
import csv
from django.db import IntegrityError
from django.http import HttpResponse
from django.contrib.auth.decorators import login_required
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer
from reportlab.lib.pagesizes import A4
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from django.db.models import Sum
from django.db.models.functions import TruncDate
from datetime import timedelta, date
import razorpay
from django.views.decorators.csrf import csrf_exempt
from django.http import HttpResponseBadRequest
from company.models import Payment
import json
from datetime import datetime
from django.db.models import Sum
from django.db.models import Q 
from django.db.models.functions import Coalesce
from inventory.models import InventoryAlert
import random
import string
import secrets
from django.core.mail import send_mail
from django.conf import settings
from django.http import JsonResponse
from django.core.validators import validate_email
from django.core.exceptions import ValidationError
import re
from notifications.services import notify_admin
from django.urls import reverse

# COMPANY OWNER LOGIN

from django.utils import timezone
from django.views.decorators.cache import never_cache

from django.template.loader import render_to_string
from django.utils.html import strip_tags
from django.core.mail import EmailMultiAlternatives
from email.mime.image import MIMEImage
from django.contrib.staticfiles import finders

import random

from django.contrib.auth.hashers import check_password

from django.core.mail import send_mail

from django.conf import settings
from django.utils import translation

from .models import PasswordHistory

@never_cache
def company_login(request):
    if request.user.is_authenticated:
        logout(request)
        messages.success(request, "You have been logged out successfully.")
        return redirect("company_login")

    if request.method == "POST":
        # Ensure we check the lowercase username if standard fails
        username = request.POST.get("username", "").strip()
        password = request.POST.get("password", "")

        user = authenticate(request, username=username, password=password)
        if not user:
            # Try lowercase username to fix case-sensitivity bug
            user = authenticate(request, username=username.lower(), password=password)

        if user:
            profile = UserProfile.objects.filter(user=user).first()

            if profile and profile.role == "COMPANY_OWNER":
                login(request, user)
                messages.success(request, f"Welcome back, {user.username}!")
                return redirect("company_dashboard")
            else:
                messages.error(request, "Access denied: This portal is for company owners only.")
        else:
            messages.error(request, "Invalid username or password.")

    return render(request, "company/company_login.html")





@never_cache
@login_required
def company_dashboard(request):
    #  Safe profile fetch
    profile = UserProfile.objects.filter(user=request.user).first()

    #  Role validation (CRITICAL FIX)
    if not profile or profile.role != "COMPANY_OWNER":
        messages.error(request, "Access denied. Please login as Company Owner.")
        return redirect("company_login")
    
    company = profile.company

    sales_orders_base_qs = SalesOrder.objects.filter(company=company)
    if company.plan_type == 'STARTER':
        sales_orders_base_qs = sales_orders_base_qs.filter(created_at__gte=timezone.now() - timedelta(days=30))

    # Welcome Banner Efficiency logic & count
    total_orders = sales_orders_base_qs.count()
    delivered_orders = sales_orders_base_qs.filter(status="DELIVERED").count()

    efficiency = 0
    if total_orders > 0:
        efficiency = round((delivered_orders/total_orders)*100)
    # 26-2-26 - Added pending shipments count to dashboard context
    pending_shipments = sales_orders_base_qs.filter(
        status__iexact='pending'
    ).count()

    # 26-2-26 Count Out of Stock Products
    products = Product.objects.filter(company=company).annotate(
        total_stock=Coalesce(Sum('productbatch__quantity'), 0)
    )

    out_of_stock_count = products.filter(total_stock=0).count()

    # Recent 5 Sales Orders with items + products (optimized)
    transactions = (
        sales_orders_base_qs
        .select_related("company")
        .prefetch_related("items__product")
        .order_by("-created_at")[:5]
    )

    # Attach product names + total quantity per order
    for order in transactions:
        order.product_names = ", ".join(
            [item.product.name for item in order.items.all()]
        )
        order.total_quantity = sum(
            item.quantity for item in order.items.all()
        )

    today = date.today()
    near_expiry_limit = today + timedelta(days=30)

    #  Product counts
    total_products = Product.objects.filter(company=company).count()

    low_stock_products = Product.objects.filter(
            company=company,
            stock_quantity__gt=0,
            stock_quantity__lte=5
        ).count()

    # 26-2-26
    # Total Revenue (only DELIVERED orders)
    total_revenue = sales_orders_base_qs.filter(
        status="DELIVERED"
    ).aggregate(total=Sum("total_amount"))["total"] or 0

    #  Batch health - Expired batches list
    expired_batches_list = ProductBatch.objects.filter(
        company=company,
        is_active=True,
        expiry_date__isnull=False,
        expiry_date__lt=date.today(),
        quantity__gt=0
    ).select_related('product').order_by('expiry_date')

    expired_batches = expired_batches_list.count()

    #  Current/Active batches list
    current_batches_list = ProductBatch.objects.filter(
        company=company,
        is_active=True,
        expiry_date__isnull=False,
        expiry_date__gte=date.today()
    ).select_related('product').order_by('expiry_date')

    near_expiry_batches = current_batches_list.filter(
        expiry_date__range=(today, near_expiry_limit),
        quantity__gt=0
    ).count()



    #  UNREAD INVENTORY ALERTS (NEW)
    unread_alerts_count = InventoryAlert.objects.filter(
        company=company,
        is_read=False
    ).count()
    # Sales Trend - Last 7 Days (DELIVERED Orders)
    last_7_days = today - timedelta(days=6)

    sales_qs = (
        SalesOrder.objects
        .filter(
            company=company,
            status="DELIVERED",
            created_at__date__gte=last_7_days
        )
        .annotate(day=TruncDate("created_at"))
        .values("day")
        .annotate(total=Sum("total_amount"))
        .order_by("day")
    )

    # Convert queryset to dictionary
    sales_dict = {item["day"]: item["total"] for item in sales_qs}

    # Ensure all 7 days appear (even if 0 sales)
    sales_trend = []
    for i in range(7):
        day = last_7_days + timedelta(days=i)
        sales_trend.append({
            "day": day.strftime("%d %b"),
            "total": float(sales_dict.get(day, 0))
        })

    # Trial Days Left / Remaining Days Logic
    days_left = None
    if company.plan_type == 'TRIAL':
        trial_end = company.created_at + timedelta(days=7)
        days_left = (trial_end.date() - today).days
        if days_left < 0:
            days_left = 0
    elif company.is_subscription_active and company.subscription_end:
        days_left = (company.subscription_end.date() - today).days
        if days_left < 0:
            days_left = 0

    import calendar
    sales_orders_used = 0
    sales_orders_limit = 100
    products_limit = 200
    days_until_reset = 0
    
    if company.plan_type == 'STARTER':
        sales_orders_used = SalesOrder.objects.filter(
            company=company,
            created_at__year=today.year,
            created_at__month=today.month
        ).count()
        _, last_day = calendar.monthrange(today.year, today.month)
        days_until_reset = last_day - today.day

    context = {
        "total_products": total_products,
        "low_stock_products": low_stock_products,
        "expired_batches": expired_batches,
        "near_expiry_batches": near_expiry_batches,
        "unread_alerts_count": unread_alerts_count,  # NEW
        "total_revenue": total_revenue, # 26-2-26
        "pending_shipments": pending_shipments, # 26-2-26
        "company": company,
        "transactions": transactions,
        "out_of_stock_count": out_of_stock_count, # 26-2-26
        "sales_trend": json.dumps(sales_trend),
        "efficiency": efficiency,
        "days_left": days_left,
        "sales_orders_used": sales_orders_used,
        "sales_orders_limit": sales_orders_limit,
        "products_limit": products_limit,
        "days_until_reset": days_until_reset,
        "transport_count": Transport.objects.filter(company=company).count() if 'Transport' in globals() else company.transports.count(),
        "pending_delete_requests": company.delete_requests.filter(status='PENDING').select_related('requested_by__user'),
        "recent_staff_activities": company.staff_activity_logs.select_related('staff__user').order_by('-timestamp')[:5],
    }

    return render(
        request,
        "company/company_dashboard.html",
        context
    )


# OPERATIONS & PAGES

from .models import NewEntry
@never_cache
@login_required
def new_entry(request):

    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    if request.method == "POST":
        entry_name = request.POST.get("entry_name")
        description = request.POST.get("description")

        if not entry_name:
            return redirect("new_entry")

        # Save entry
        NewEntry.objects.create(
            company=company,
            entry_name=entry_name,
            description=description
        )

        # Redirect with success flag
        return redirect("/company/new-entry/?success=1")

    entries = NewEntry.objects.filter(company=company).order_by("-created_at")

    context = {
        "company": company,
        "entries": entries
    }

    return render(request, "company/new_entry.html", context)

@login_required
def delete_entry(request, entry_id):

    entry = NewEntry.objects.get(id=entry_id)

    entry.delete()

    return redirect("new_entry")

def reports_page(request):
    profile = getattr(request.user, 'userprofile', None)
    company = profile.company if profile else None
    if profile and company.plan_type == 'STARTER':
        messages.error(request, 'Upgrade to Business plan to unlock all reports.')
        return redirect('company_dashboard')
    return render(request, "company/reports.html", {'company': company})  # ← fixed

@login_required
def add_stock(request):
    return render(request, "company/add_stock.html")

@login_required
def dispatcher(request):
    profile = UserProfile.objects.select_related("company").get(user=request.user)
    if profile.company.plan_type == 'STARTER':
        messages.error(request, "FEFO Dispatch logic is restricted to the Business Plan.")
        return redirect('company_dashboard')

    if request.method == "POST":
        order_number = request.POST.get('order_number')
        dispatch_date = request.POST.get('dispatch_date')

        if not dispatch_date:
            return render(request, 'company/dispatch.html', {
                'error': 'Dispatch Date is required.'
            })

        messages.success(request, "Order dispatched successfully.")
        return redirect('dispatch')
    return render(request, 'company/dispatch.html', {'company': profile.company})

# ── STAFF CREDENTIAL HELPERS ─────────────────────────────────────────────────

def _generate_staff_username(name, company_name):
    """
    Build a unique Django username: <firstname>.<company_slug>
    e.g. "rahul.acmecorp". Appends a counter if already taken.
    """
    from django.contrib.auth.models import User
    first = name.strip().split()[0].lower()
    co_slug = re.sub(r'[^a-z0-9]', '', company_name.lower())[:8]
    base = f"{first}.{co_slug}"[:28]
    username = base
    counter = 1
    while User.objects.filter(username=username).exists():
        username = f"{base}{counter}"
        counter += 1
    return username


def _generate_password(length=10):
    """
    Return a cryptographically random password containing at least one
    uppercase letter, one lowercase letter, one digit, and one special char.
    """
    alphabet = string.ascii_letters + string.digits + "!@#$%&"
    while True:
        pwd = ''.join(secrets.choice(alphabet) for _ in range(length))
        if (any(c.isupper() for c in pwd) and
                any(c.islower() for c in pwd) and
                any(c.isdigit() for c in pwd) and
                any(c in "!@#$%&" for c in pwd)):
            return pwd


def _send_staff_credentials_email(request, staff, plain_password, company):
    """Send an HTML email to the staff member with their login credentials."""
    login_url = request.build_absolute_uri('/staff/login/')

    html_message = render_to_string(
        'company/emails/staff_credentials_email.html',
        {
            'staff_name': staff.name,
            'company_name': company.name,
            'username': staff.username,
            'password': plain_password,
            'login_url': login_url,
        }
    )
    plain_message = strip_tags(html_message)

    email_msg = EmailMultiAlternatives(
        subject=f"Your Staff Portal Credentials — {company.name}",
        body=plain_message,
        from_email=settings.EMAIL_HOST_USER,
        to=[staff.email],
    )
    email_msg.attach_alternative(html_message, "text/html")

    logo_path = finders.find('accounts/images/logo.png')
    if logo_path:
        from email.mime.image import MIMEImage
        with open(logo_path, 'rb') as f:
            mime_image = MIMEImage(f.read())
            mime_image.add_header('Content-ID', '<logo>')
            mime_image.add_header('Content-Disposition', 'inline', filename='logo.png')
            email_msg.attach(mime_image)

    email_msg.send(fail_silently=False)


# ── STAFF INVITE EMAIL HELPER ────────────────────────────────────────────────

def _send_staff_invite_email(request, staff, company):
    """Send an invite link so the staff member sets their own password."""
    from django.urls import reverse
    invite_url = request.build_absolute_uri(
        reverse('staff_accept_invite', kwargs={'token': str(staff.invite_token)})
    )
    html_message = render_to_string(
        'company/emails/staff_invite_email.html',
        {
            'staff_name': staff.name,
            'company_name': company.name,
            'invite_url': invite_url,
        }
    )
    plain_message = strip_tags(html_message)
    email_msg = EmailMultiAlternatives(
        subject=f"You're invited to the Staff Portal — {company.name}",
        body=plain_message,
        from_email=settings.EMAIL_HOST_USER,
        to=[staff.email],
    )
    email_msg.attach_alternative(html_message, "text/html")
    logo_path = finders.find('accounts/images/logo.png')
    if logo_path:
        from email.mime.image import MIMEImage
        with open(logo_path, 'rb') as f2:
            mime_image = MIMEImage(f2.read())
            mime_image.add_header('Content-ID', '<logo>')
            mime_image.add_header('Content-Disposition', 'inline', filename='logo.png')
            email_msg.attach(mime_image)
    email_msg.send(fail_silently=False)


# ── ADD STAFF ─────────────────────────────────────────────────────────────────

@login_required
def add_staff(request):
    from django.contrib.auth.models import User

    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    if request.method == "POST":
        name  = request.POST.get("staff_name", "").strip()
        email = request.POST.get("email", "").strip()

        if not name:
            messages.error(request, "Staff name is required.")
            return render(request, "company/add_staff.html")

        if not email:
            messages.error(request, "Email is required.")
            return render(request, "company/add_staff.html")

        if Staff.objects.filter(email=email).exists():
            messages.error(request, "A staff member with this email already exists.")
            return render(request, "company/add_staff.html")

        # Generate credentials
        username       = _generate_staff_username(name, company.name)
        plain_password = _generate_password()

        # Create Django User
        user = User.objects.create_user(
            username=username,
            email=email,
            password=plain_password,
            first_name=name.split()[0],
            last_name=' '.join(name.split()[1:]) if len(name.split()) > 1 else '',
        )

        # Create Staff record linked to that User
        staff = Staff.objects.create(
            company=company,
            name=name,
            email=email,
            user=user,
            username=username,
        )

        # Email credentials to staff member
        try:
            _send_staff_credentials_email(request, staff, plain_password, company)
            messages.success(
                request,
                f"Staff member '{name}' added. "
                f"Login credentials emailed to {email}."
            )
        except Exception as e:
            messages.warning(
                request,
                f"Staff '{name}' added, but credential email failed ({e}). "
                f"Share manually — Username: {username}"
            )

        return redirect("staff_list")

    return render(request, "company/add_staff.html")


# ── DELETE STAFF ──────────────────────────────────────────────────────────────

@login_required
def delete_staff(request, staff_id):
    """Remove a staff member and their linked Django User account."""
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    staff = get_object_or_404(Staff, pk=staff_id, company=company)

    if request.method == "POST":
        staff_name = staff.name
        linked_user = staff.user
        staff.delete()
        if linked_user:
            linked_user.delete()
        messages.success(request, f"Staff member '{staff_name}' has been removed.")

    return redirect("staff_list")



# ── RESEND STAFF INVITE ───────────────────────────────────────────────────────

@login_required
def resend_staff_invite(request, staff_id):
    """Generate a fresh invite token and re-send the invite link email."""
    import uuid
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company
    staff = get_object_or_404(Staff, pk=staff_id, company=company)

    # Only resend if not yet accepted
    if staff.invite_accepted:
        messages.warning(request, f"{staff.name} has already accepted their invite.")
        return redirect("staff_list")

    # Rotate the token so old links stop working
    staff.invite_token = uuid.uuid4()
    staff.save(update_fields=['invite_token'])

    try:
        _send_staff_invite_email(request, staff, company)
        messages.success(request, f"Invite re-sent to {staff.email}.")
    except Exception as e:
        messages.warning(
            request,
            f"Couldn't send email ({e}). Share this link manually: "
            f"/staff/invite/{staff.invite_token}/"
        )
    return redirect("staff_list")


# ── REGENERATE STAFF CREDENTIALS ─────────────────────────────────────────────

@login_required
def regenerate_staff_credentials(request, staff_id):
    """Reset a staff member's password and email the new credentials."""
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    staff = get_object_or_404(Staff, pk=staff_id, company=company)

    if request.method == "POST":
        if not staff.user:
            messages.error(request, "This staff member has no linked user account.")
            return redirect("staff_list")

        plain_password = _generate_password()
        staff.user.set_password(plain_password)
        staff.user.save()

        try:
            _send_staff_credentials_email(request, staff, plain_password, company)
            messages.success(
                request,
                f"New credentials generated and emailed to {staff.email}."
            )
        except Exception as e:
            messages.warning(
                request,
                f"Password reset, but email failed ({e}). "
                f"New password: {plain_password}"
            )

    return redirect("staff_list")

@login_required
def export_data(request):
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    # When user clicks Export Now
    if request.method == "POST":
        return redirect("export_data_download")

    return render(request, "company/export.html", {"company": company})




@login_required
def export_data_download(request):
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    response = HttpResponse(content_type="text/csv")
    response["Content-Disposition"] = 'attachment; filename="company_data.csv"'

    writer = csv.writer(response)
    writer.writerow(["Product Name", "Stock Quantity"])

    products = Product.objects.filter(company=company)

    for p in products:
        writer.writerow([p.name, p.stock_quantity])

    return response

@login_required
def transactions_list(request):
    profile = UserProfile.objects.select_related("company").filter(user=request.user).first()

    if not profile or profile.role != "COMPANY_OWNER":
        return redirect("company_login")

    # Get all transactions
    transaction_list = Transaction.objects.filter(company=profile.company).order_by("-created_at")
    
    # Set up pagination (e.g., 10 transactions per page)
    paginator = Paginator(transaction_list, 10) 
    page_number = request.GET.get('page')
    transactions = paginator.get_page(page_number)

    return render(request, "company/transactions_list.html", {"transactions": transactions})


# COMPANY OWNER LOGOUT

def company_logout(request):
    user_role = None
    is_superuser = False

    if request.user.is_authenticated:
        is_superuser = request.user.is_superuser  #  store BEFORE logout
        
        try:
            profile = UserProfile.objects.get(user=request.user)
            user_role = profile.role
        except UserProfile.DoesNotExist:
            pass

    logout(request)

    messages.success(request, "You have been logged out successfully.")

    #  Now use stored values
    if user_role == 'COMPANY_OWNER':
        return redirect('company_login')
    elif is_superuser:
        return redirect('admin_login')

    return redirect('company_login')


# SETTINGS

# SETTINGS
# SETTINGS
@login_required
def company_settings(request):
    profile = UserProfile.objects.select_related("company").get(user=request.user)

    if profile.role != "COMPANY_OWNER":
        return redirect("company_login")

    company = profile.company

    if request.method == "POST":
        name = request.POST.get("company_name", "").strip()
        email = request.POST.get("company_email", "").strip()
        new_password = request.POST.get("new_password")
        confirm_password = request.POST.get("confirm_password")
        phone = request.POST.get("company_phone", "").strip()
        address = request.POST.get("company_address", "").strip()
        gst_number = request.POST.get("gst_number", "").strip().upper()
        logo = request.FILES.get("company_logo")

        errors = {}

        # Name validation
        if not name:
            errors["name_error"] = "Company name is required."
        elif len(name) > 100:
            errors["name_error"] = "Company name cannot exceed 100 characters."

        # Email validation
        EMAIL_REGEX = r'^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$'
        if not email:
            errors["email_error"] = "Email is required."
        elif not re.match(EMAIL_REGEX, email):
            errors["email_error"] = "Enter a valid email address."

        # Phone validation — Indian mobile numbers only (start with 6-9)
        PHONE_REGEX = r'^[6-9]\d{9}$'
        if not phone:
            errors["phone_error"] = "Phone number is required."
        elif not phone.isdigit():
            errors["phone_error"] = "Phone number must contain digits only."
        elif not re.match(PHONE_REGEX, phone):
            errors["phone_error"] = "Enter a valid 10-digit number starting with 6, 7, 8 or 9."

        # GST validation (optional field)
        GST_REGEX = r'^[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z]{1}[A-Z0-9]{3}$'
        if gst_number and not re.match(GST_REGEX, gst_number):
            errors["gst_error"] = "Enter a valid 15-character GST number (e.g. 22AAAAA0000A1Z5)."

        # If any errors — return to form with errors and user's entered data
        if errors:
            company.name = name
            company.email = email
            company.phone = phone
            company.address = address
            company.gst_number = gst_number
            return render(request, "company/company_settings.html", {
                "company": company,
                "currency_choices": company.CURRENCY_CHOICES,
                **errors,
            })

        company.name = name
        company.email = email
        company.phone = phone
        company.address = address
        company.gst_number = gst_number
        if logo:
            company.logo = logo

        currency = request.POST.get("company_currency")
        if currency:
            company.display_currency = currency
            
        language = request.POST.get("company_language")
        if language:
            company.language = language

        try:
            company.full_clean()  # Validate model fields
            company.save()
        except ValidationError as e:

            return render(
                request,
                "company/company_settings.html",
                {
                    "company": company,
                    "currency_choices": company.CURRENCY_CHOICES,
                    "gst_error": e.message_dict.get("gst_number", [None])[0],
                }
            )
        except IntegrityError:
            return render(request, "company/company_settings.html", {
                "company": company,
                "currency_choices": company.CURRENCY_CHOICES,
                "email_error": "This email is already registered by another company.",
            })

        if new_password:
            if new_password != confirm_password:
                messages.error(request, "Passwords do not match")
                return redirect("company_settings")

            otp_verified = request.session.get(
                "otp_verified",
                 False
            )

            if not otp_verified:
                messages.error(
                    request,
                    "Please verify OTP first"
                )
                return redirect(
                    "company_settings"
                )

            if len(new_password) < 8:
                messages.error(request, "Password must be at least 8 characters")
                return redirect("company_settings")

            if not any(c.isupper() for c in new_password):
                messages.error(request, "Password must contain at least one uppercase letter")
                return redirect("company_settings")

            if not any(c.islower() for c in new_password):
                messages.error(request, "Password must contain at least one lowercase letter")
                return redirect("company_settings")

            if not any(c.isdigit() for c in new_password):
                messages.error(request, "Password must contain at least one number")
                return redirect("company_settings")

            if not any(c in "!@#$%&" for c in new_password):
                messages.error(request, "Password must contain at least one special character")
                return redirect("company_settings")

            last_passwords = PasswordHistory.objects.filter(
                user=request.user
            ).order_by("-created_at")[:3]

            for item in last_passwords:

                if check_password(
                    new_password,
                    item.password
                ):
                    messages.error(
                        request,
                        "Last 3 passwords cannot be reused"
                    )

                    return redirect(
                        "company_settings"
                    )

            current_user = request.user

            current_user.set_password(
                new_password
            )

            current_user.save()

            PasswordHistory.objects.create(
                user=current_user,
                password=current_user.password
            )

            update_session_auth_hash(
                request,
                current_user
            )

            send_mail(
                "Password Changed Successfully",
                "Your password has been changed successfully.",
                settings.EMAIL_HOST_USER,
                [request.user.email],
                fail_silently=False
            )

            request.session["otp_verified"] = False

        response = redirect("company_settings")
        if language:
            translation.activate(language)
            response.set_cookie(settings.LANGUAGE_COOKIE_NAME, language)
            
        messages.success(request, "Settings updated successfully")
        return response

    return render(request, "company/company_settings.html", {
        "company": company,
        "currency_choices": company.CURRENCY_CHOICES,
    })




@login_required
@csrf_exempt
def verify_old_password(request):
    if request.method == "POST":
        try:
            data = json.loads(request.body)
            old_password = data.get("old_password", "")
            if request.user.check_password(old_password):
                return JsonResponse({"success": True})
            return JsonResponse({"success": False, "message": "Incorrect old password."})
        except Exception as e:
            return JsonResponse({"success": False, "message": str(e)})
    return JsonResponse({"success": False, "message": "Invalid request"})



@login_required
def export_stock_movement_pdf(request, pk=None):
    from company.pdf_fonts import get_unicode_font, get_unicode_font_bold
    ufont = get_unicode_font()
    ufont_bold = get_unicode_font_bold()

    company = Company.objects.get(owner=request.user)

    # Base queryset
    transactions = StockTransaction.objects.filter(
        product__company=company
    )

    product = None

    # Filter by product (optional)
    if pk:
        product = get_object_or_404(Product, id=pk, company=company)
        transactions = transactions.filter(product=product)

    # Apply filters
    from_date = request.GET.get("from_date")
    to_date = request.GET.get("to_date")
    flow_type = request.GET.get("flow_type")

    if from_date:
        transactions = transactions.filter(created_at__date__gte=from_date)

    if to_date:
        transactions = transactions.filter(created_at__date__lte=to_date)

    if flow_type:
        transactions = transactions.filter(transaction_type=flow_type)

    transactions = transactions.order_by('-created_at')

    # Response
    response = HttpResponse(content_type='application/pdf')

    filename = f"{company.name}_stock_report.pdf"
    if product:
        filename = f"{product.name}_stock_report.pdf"

    response['Content-Disposition'] = f'attachment; filename="{filename}"'

    # PDF setup
    doc = SimpleDocTemplate(
        response,
        pagesize=A4,
        rightMargin=30,
        leftMargin=30,
        topMargin=40,
        bottomMargin=20
    )

    elements = []
    styles = getSampleStyleSheet()

    # Custom styles
    header_style = ParagraphStyle(
        "Header",
        parent=styles["Heading1"],
        fontSize=16,
        fontName=ufont_bold,
        textColor=colors.white,
        alignment=1,
        spaceAfter=6
    )

    table_text_style = ParagraphStyle(
        name="TableText",
        fontName=ufont,
        fontSize=9,
        leading=11,
    )

    # Header title
    title = f"{company.name} - Stock Movement Report"
    if product:
        title = f"{product.name} - Stock Movement Report"

    header_table = Table([[Paragraph(title, header_style)]], colWidths=[520])
    header_table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#4071db")),
        ("LEFTPADDING", (0, 0), (-1, -1), 10),
        ("RIGHTPADDING", (0, 0), (-1, -1), 10),
        ("TOPPADDING", (0, 0), (-1, -1), 6),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 6),
    ]))

    elements.append(header_table)
    elements.append(Spacer(1, 8))

    # Generated date
    date_style = ParagraphStyle(name="DateStyle", parent=styles["Normal"], fontName=ufont)
    elements.append(Paragraph(
        f"Generated on: {datetime.now().strftime('%d %b %Y, %I:%M %p')}",
        date_style
    ))
    elements.append(Spacer(1, 15))

    # Table data
    data = [["Date", "Product", "Type", "Quantity"]]

    for t in transactions:
        data.append([
            t.created_at.strftime("%d-%m-%Y"),
            Paragraph(t.product.name, table_text_style),  #  wrapped text
            Paragraph(t.transaction_type, table_text_style),
            Paragraph(str(t.quantity), table_text_style),
        ])

    # Table with better column widths
    table = Table(data, colWidths=[100, 240, 90, 70])

    table.setStyle(TableStyle([
        # Header styling
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f172a")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), ufont_bold),
        ("FONTSIZE", (0, 0), (-1, 0), 10),

        # Body styling
        ("FONTNAME", (0, 1), (-1, -1), ufont),
        ("FONTSIZE", (0, 1), (-1, -1), 9),
        ("ALIGN", (3, 1), (3, -1), "CENTER"),

        # Alternating rows
        ("ROWBACKGROUNDS", (0, 1), (-1, -1),
         [colors.HexColor("#f1f5f9"), colors.white]),

        # Grid
        ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),

        # Padding
        ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
        ("TOPPADDING", (0, 0), (-1, -1), 8),

        # Wrap product column
        ('WORDWRAP', (1, 1), (1, -1), 'CJK'),
    ]))

    elements.append(table)

    doc.build(elements)
    return response

@login_required
def staff_list(request):
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    staff_members = Staff.objects.filter(company=company).select_related('user')
    return render(request, "company/staff_list.html", {
        "staff_members": staff_members,
        "company": company,
    })

def forgot_password(request):
    if request.user.is_authenticated:
        profile = UserProfile.objects.filter(user=request.user).first()

        if profile and profile.role == "COMPANY_OWNER":
            return redirect("company_dashboard")

    if request.method == "POST":
        email = request.POST.get("email")
        profile = Company.objects.filter(email=email).first()

        if not profile:
            messages.error(request,"No account found with this email")
            return redirect("forgot_password")

        otp = random.randint(100000, 999999)

        request.session['reset_otp'] = str(otp)
        request.session['reset_user'] = profile.id
        request.session["otp_resend_count"] = 0

        subject = "Password Reset OTP | CWS Inventory"
        html_message = render_to_string('company/emails/reset_password_otp.html',{
                'name': profile.name,
                'otp': otp,
            })

        plain_message = strip_tags(html_message)

        email_msg = EmailMultiAlternatives(subject,plain_message,settings.EMAIL_HOST_USER,[email])
        email_msg.attach_alternative(html_message,"text/html")
        logo_path = finders.find('accounts/images/logo.png')

        if logo_path:
            with open(logo_path, 'rb') as f:
                mime_image = MIMEImage(f.read())
                mime_image.add_header('Content-ID','<logo>')
                mime_image.add_header('Content-Disposition','inline',filename='logo.png')
                email_msg.attach(mime_image)

        email_msg.send(fail_silently=False)
        return redirect("verify_otp")

    return render(request,"company/forgot_password.html")


def verify_otp(request):

    if request.method == "POST":

        entered_otp = request.POST.get("otp")
        session_otp = request.session.get("reset_otp")

        if str(entered_otp) == str(session_otp):

            return redirect("reset_password")

        else:
            messages.error(request,"Invalid OTP")
            return redirect("forgot_password")

    return render(request,"company/verify_otp.html")



def resend_otp(request):
    resend_count = request.session.get("otp_resend_count",0)

    if resend_count >= 3:
        return JsonResponse({"error": "limit reached"})

    user_id = request.session.get("reset_user")
    user = Company.objects.get(id=user_id)
    otp = random.randint(100000,999999)
    request.session["reset_otp"] = str(otp)
    request.session["otp_resend_count"] = (resend_count + 1)

    subject = ("New Password Reset OTP | CWS Inventory")
    html_message = render_to_string('accounts/emails/reset_password_otp.html',{
        'name': user.name,
            'otp': otp,
        })

    plain_message = strip_tags(html_message)
    email_msg = EmailMultiAlternatives(subject,plain_message,settings.EMAIL_HOST_USER,[user.email])
    email_msg.attach_alternative(html_message,"text/html")
    logo_path = finders.find('accounts/images/logo.png')

    if logo_path:
        with open(logo_path, 'rb') as f:
            mime_image = MIMEImage(f.read())
            mime_image.add_header('Content-ID','<logo>')
            mime_image.add_header('Content-Disposition','inline',filename='logo.png')
            email_msg.attach(mime_image)

    email_msg.send(fail_silently=False)
    return JsonResponse({"status": "sent"})


def reset_password(request):

    user_id = request.session.get("reset_user")

    if not user_id:
        messages.error(request, "Session expired. Please try again.")
        return redirect("forgot_password")

    if request.method == "POST":

        password = request.POST.get("password")
        confirm_password = request.POST.get("confirm_password")

        if password != confirm_password:
            messages.error(request, "Passwords do not match")
            return redirect("reset_password")

        company = Company.objects.get(id=user_id)
        user = company.owner

        user.set_password(password)
        user.save()

        request.session.pop("reset_user", None)
        request.session.pop("reset_otp", None)

        return redirect("company_login")

    return render(request, "company/reset_password.html")




# RAZORPAY CLIENT
client = razorpay.Client(
    auth=(
        settings.RAZORPAY_KEY_ID,
        settings.RAZORPAY_KEY_SECRET
    )
)


@login_required
def pricing(request):
    return render(request,'company/pricing.html')


@csrf_exempt
@login_required
def create_order(request):
    if request.method != "POST":
        return JsonResponse({
            "status": "failed",
            "message": "Invalid request method"
        })

    try:
        data = json.loads(request.body)
        plan = data.get("plan")

        VALID_PLANS = {"STARTER": 499,"BUSINESS": 1299}

        if plan not in VALID_PLANS:
            return JsonResponse({
                "status": "failed",
                "message": "Invalid Plan Selected"
            })

        try:
            company = Company.objects.get(owner=request.user)

        except Company.DoesNotExist:
            return JsonResponse({"status": "failed","message": "Company not found"})

        amount = VALID_PLANS[plan] * 100

        order_data = {
            "amount": amount,
            "currency": "INR",
            "payment_capture": 1,

            "notes": {
                "company_id": str(company.id),
                "company_name": company.name,
                "plan_name": plan
            }
        }

        payment_order = client.order.create(order_data)

        return JsonResponse({
            "status": "success",
            "order_id": payment_order["id"],
            "amount": payment_order["amount"],
            "razorpay_key": settings.RAZORPAY_KEY_ID,
            "plan": plan
        })

    except Exception as e:
        print(e)
        return JsonResponse({
            "status": "failed",
            "message": "Something went wrong"
        })


@csrf_exempt
def payment_success(request):

    if request.method != "POST":
        return HttpResponseBadRequest("Invalid Request Method")

    try:
        payment_id = request.POST.get('razorpay_payment_id')
        order_id = request.POST.get('razorpay_order_id')
        signature = request.POST.get('razorpay_signature')

        params_dict = {
            'razorpay_payment_id': payment_id,
            'razorpay_order_id': order_id,
            'razorpay_signature': signature
        }

        client.utility.verify_payment_signature(params_dict)
        payment_data = client.payment.fetch(payment_id)

        if payment_data['status'] != "captured":
            return redirect( 'payment_failed' )

        notes = payment_data.get("notes",{})
        company_id = notes.get("company_id")
        plan_name = notes.get("plan_name")

        try:
            company = Company.objects.get(id=company_id)

        except Company.DoesNotExist:
            return HttpResponseBadRequest("Company Not Found")

        Payment.objects.create(
            company=company,
            plan_name=plan_name,
            order_id=order_id,
            payment_id=payment_id,
            razorpay_signature=signature,
            payment_status=payment_data['status'],
            amount=payment_data['amount'] / 100,
            currency=payment_data['currency'],
            is_verified=True
        )

        # Upgrade subscription logic
        if company.is_subscription_active and company.subscription_end and company.subscription_end > timezone.now():
            # Add 30 days to the existing subscription end date
            company.subscription_end = company.subscription_end + timedelta(days=30)
        else:
            # Start a new 30-day subscription from today
            company.subscription_start = timezone.now()
            company.subscription_end = timezone.now() + timedelta(days=30)

        company.plan_type = plan_name
        company.is_subscription_active = True
        company.save()

        notify_admin(
            'Subscription',
            'Subscription Renewal',
            f'Company <strong>{company.name}</strong> renewed subscription for plan <strong>{plan_name}</strong>.',
            '/admin/companies/'
        )

        notify_admin(
            'Subscription',
            'Payment successful',
            f'Payment received from <strong>{company.name}</strong>.<br>Amount: {payment_data["amount"] / 100} {payment_data["currency"]}.',
            reverse('company_list')
        )

        return render(request,'company/payment_success.html',{
            'payment_id': payment_id, 
            'company': company, 
            'plan_name': plan_name, 
            'amount': payment_data['amount'] // 100
        })

    except Exception as e:
        print(e)
        return redirect( 'payment_failed' )
    

@login_required
def payment_failed(request):

    company = None
    try:
        profile = request.user.userprofile
        company = profile.company
    except:
        pass

    if company:
        notify_admin(
            'Subscription',
            'Payment failed',
            f'Payment failed for company <strong>{company.name}</strong>.',
            '/admin/companies/'
        )
    return render(request,'company/payment_failed.html')




from .models import Transport
from .forms import TransportForm



@login_required
def add_transport(request):

    if request.user.is_superuser:
        company = Company.objects.first()
    else:
        profile = UserProfile.objects.filter(
            user=request.user
        ).first()

        if not profile:
            messages.error(
                request,
                "User Profile not found"
            )
            return redirect('admin_dashboard')

        company = profile.company

    if request.method == "POST":
        form = TransportForm(request.POST)

        if form.is_valid():
            transport = form.save(commit=False)
            transport.company = company
            transport.save()

            return redirect("transport_list")

    else:
        form = TransportForm()

    return render(
        request,
        "company/add_transport.html",
        {"form": form,
        "company": company}
    )






from .models import Transport

@login_required
def transport_list(request):
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    transports = Transport.objects.filter(
        company=company
    ).order_by("-id")

    return render(
        request,
        "company/transport_list.html",
        {
            "transports": transports,
            "company": company
        }
    )


from django.shortcuts import get_object_or_404

@login_required
def edit_transport(request, pk):

    profile = UserProfile.objects.get(
        user=request.user
    )

    company = profile.company

    transport = get_object_or_404(
        Transport,
        pk=pk,
        company=company
    )

    if request.method == "POST":
        form = TransportForm(
            request.POST,
            instance=transport
        )

        if form.is_valid():
            form.save()
            return redirect('transport_list')

    else:
        form = TransportForm(instance=transport)

    return render(
        request,
        "company/edit_transport.html",
        {
            'form': form,
            'transport': transport,
            'company': company,
        }
    )



from django.shortcuts import get_object_or_404, redirect

@login_required
def delete_transport(request, pk):

    profile = UserProfile.objects.get(
        user=request.user
    )

    company = profile.company

    transport = get_object_or_404(
        Transport,
        id=pk,
        company=company
    )

    if request.method == "POST":

        transport.delete()

        messages.success(
            request,
            "Transport deleted successfully."
        )

    return redirect("transport_list")


from company.models import Transport

@login_required
def admin_dashboard(request):

    total_transport = Transport.objects.count()

    pending_transport = Transport.objects.filter(
        status='PENDING'
    ).count()

    in_transit_transport = Transport.objects.filter(
        status='IN_TRANSIT'
    ).count()

    delivered_transport = Transport.objects.filter(
        status='DELIVERED'
    ).count()

    context = {
        'total_transport': total_transport,
        'pending_transport': pending_transport,
        'in_transit_transport': in_transit_transport,
        'delivered_transport': delivered_transport,
    }

    return render(
        request,
        'accounts/admin_dashboard.html',
        context
    )


# ─────────────────────────────────────────────────────────────────
# TRANSPORT — EXPORT EXCEL
# ─────────────────────────────────────────────────────────────────
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

@login_required
def export_transport_excel(request):
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    transports = Transport.objects.filter(company=company).order_by("-id")

    wb = Workbook()
    ws = wb.active
    ws.title = "Transport Report"

    # ── Column widths ──────────────────────────────────────────
    col_widths = [8, 18, 20, 16, 14, 22, 22, 16, 22, 14]
    for i, w in enumerate(col_widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w

    # ── Header row styling ─────────────────────────────────────
    header_fill   = PatternFill("solid", fgColor="1A5276")
    header_font   = Font(bold=True, color="FFFFFF", size=10)
    center_align  = Alignment(horizontal="center", vertical="center", wrap_text=True)
    thin_border   = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin"),
    )

    headers = [
        "#", "Vehicle No", "Driver Name", "Driver Mobile",
        "Type", "Source", "Destination",
        "Dispatch Date", "Expected Delivery", "Status"
    ]

    ws.append(headers)
    for col_idx, _ in enumerate(headers, start=1):
        cell = ws.cell(1, col_idx)
        cell.font      = header_font
        cell.fill      = header_fill
        cell.alignment = center_align
        cell.border    = thin_border

    ws.row_dimensions[1].height = 30

    # ── Status colour map ──────────────────────────────────────
    STATUS_FILL = {
        "PENDING":    PatternFill("solid", fgColor="FFF3CD"),
        "IN_TRANSIT": PatternFill("solid", fgColor="CCE5FF"),
        "DELIVERED":  PatternFill("solid", fgColor="D4EDDA"),
        "CANCELLED":  PatternFill("solid", fgColor="F8D7DA"),
    }
    alt_fill = PatternFill("solid", fgColor="EBF5FB")

    # ── Data rows ──────────────────────────────────────────────
    for idx, t in enumerate(transports, start=1):
        row = [
            idx,
            t.vehicle_number,
            t.driver_name,
            t.driver_mobile,
            t.get_transport_type_display(),
            t.source_location,
            t.destination_location,
            t.dispatch_date.strftime("%d-%m-%Y") if t.dispatch_date else "-",
            t.expected_delivery_date.strftime("%d-%m-%Y") if t.expected_delivery_date else "-",
            t.get_status_display(),
        ]
        ws.append(row)

        excel_row = idx + 1   # row 1 is header
        row_fill  = STATUS_FILL.get(t.status, alt_fill if idx % 2 == 0 else None)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(excel_row, col_idx)
            cell.border    = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            # apply status colour only to Status cell (col 10), alt shade rest
            if col_idx == 10:
                cell.fill = STATUS_FILL.get(t.status, PatternFill())
            elif idx % 2 == 0:
                cell.fill = alt_fill

    # ── Freeze header row ──────────────────────────────────────
    ws.freeze_panes = "A2"

    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = (
        f'attachment; filename="{company.name}_transport_report.xlsx"'
    )
    wb.save(response)
    return response


# ─────────────────────────────────────────────────────────────────
# TRANSPORT — EXPORT PDF
# ─────────────────────────────────────────────────────────────────
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib import colors
from reportlab.lib.units import mm
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib.enums import TA_LEFT, TA_CENTER
from datetime import datetime as dt

@login_required
def export_transport_pdf(request):
    from company.pdf_fonts import get_unicode_font, get_unicode_font_bold
    ufont = get_unicode_font()
    ufont_bold = get_unicode_font_bold()

    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    transports = Transport.objects.filter(company=company).order_by("-id")

    response = HttpResponse(content_type="application/pdf")
    response["Content-Disposition"] = (
        f'attachment; filename="{company.name}_transport_report.pdf"'
    )

    # landscape A4 — more columns fit comfortably
    doc = SimpleDocTemplate(
        response,
        pagesize=landscape(A4),
        leftMargin=15*mm, rightMargin=15*mm,
        topMargin=12*mm,  bottomMargin=12*mm,
    )

    styles    = getSampleStyleSheet()
    elements  = []

    # ── Styles ─────────────────────────────────────────────────
    title_style = ParagraphStyle(
        "title", parent=styles["Heading1"],
        fontName=ufont_bold,
        fontSize=15, textColor=colors.white,
        alignment=1, spaceAfter=4,
    )
    cell_style = ParagraphStyle(
        "cell", parent=styles["Normal"],
        fontName=ufont,
        fontSize=8, leading=10, wordWrap="CJK",
    )
    center_cell = ParagraphStyle(
        "ccell", parent=styles["Normal"],
        fontName=ufont,
        fontSize=8, leading=10, alignment=1,
    )

    # ── Banner ─────────────────────────────────────────────────
    banner = Table(
        [[Paragraph(f"{company.name} — Transport Management Report", title_style)]],
        colWidths=[267*mm],
    )
    banner.setStyle(TableStyle([
        ("BACKGROUND",    (0,0),(-1,-1), colors.HexColor("#1A5276")),
        ("TOPPADDING",    (0,0),(-1,-1), 8),
        ("BOTTOMPADDING", (0,0),(-1,-1), 8),
    ]))
    elements.append(banner)
    elements.append(Spacer(1, 4*mm))
    sub_style = ParagraphStyle("sub", parent=styles["Normal"], fontName=ufont, fontSize=8,
                               textColor=colors.grey)
    elements.append(Paragraph(
        f"Generated: {dt.now().strftime('%d %b %Y, %I:%M %p')}  |  "
        f"Total Records: {transports.count()}",
        sub_style,
    ))
    elements.append(Spacer(1, 4*mm))

    # ── Status colour map ──────────────────────────────────────
    STATUS_BG = {
        "PENDING":    colors.HexColor("#FFF3CD"),
        "IN_TRANSIT": colors.HexColor("#CCE5FF"),
        "DELIVERED":  colors.HexColor("#D4EDDA"),
        "CANCELLED":  colors.HexColor("#F8D7DA"),
    }

    # ── Table data ─────────────────────────────────────────────
    # landscape A4 content width ≈ 267mm; columns must sum to that
    col_widths = [10*mm, 28*mm, 32*mm, 28*mm, 20*mm, 38*mm, 38*mm, 26*mm, 26*mm, 21*mm]

    data = [[
        "#", "Vehicle No", "Driver Name", "Mobile",
        "Type", "Source", "Destination",
        "Dispatch", "Exp. Delivery", "Status"
    ]]

    for idx, t in enumerate(transports, start=1):
        data.append([
            str(idx),
            Paragraph(t.vehicle_number, cell_style),
            Paragraph(t.driver_name, cell_style),
            t.driver_mobile,
            t.get_transport_type_display(),
            Paragraph(t.source_location, cell_style),
            Paragraph(t.destination_location, cell_style),
            t.dispatch_date.strftime("%d-%m-%Y") if t.dispatch_date else "-",
            t.expected_delivery_date.strftime("%d-%m-%Y") if t.expected_delivery_date else "-",
            Paragraph(t.get_status_display(), center_cell),
        ])

    table = Table(data, colWidths=col_widths, repeatRows=1)

    # base style
    ts = [
        # Header
        ("BACKGROUND",    (0,0),(-1,0),  colors.HexColor("#1A5276")),
        ("TEXTCOLOR",     (0,0),(-1,0),  colors.white),
        ("FONTNAME",      (0,0),(-1,0),  ufont_bold),
        ("FONTSIZE",      (0,0),(-1,0),  8),
        ("ALIGN",         (0,0),(-1,0),  "CENTER"),
        ("BOTTOMPADDING", (0,0),(-1,0),  6),
        ("TOPPADDING",    (0,0),(-1,0),  6),
        # Data
        ("FONTNAME",      (0,1),(-1,-1), ufont),
        ("FONTSIZE",      (0,1),(-1,-1), 8),
        ("VALIGN",        (0,0),(-1,-1), "MIDDLE"),
        ("ALIGN",         (0,1),(0,-1),  "CENTER"),   # # col
        ("TOPPADDING",    (0,1),(-1,-1), 4),
        ("BOTTOMPADDING", (0,1),(-1,-1), 4),
        # Grid
        ("GRID",          (0,0),(-1,-1), 0.4, colors.HexColor("#AED6F1")),
        ("LINEBELOW",     (0,0),(-1,0),  1,   colors.HexColor("#1A5276")),
        # Alternating rows
        *[("BACKGROUND",  (0,r),(-1,r),  colors.HexColor("#EBF5FB"))
          for r in range(2, len(data), 2)],
    ]

    # per-row status colour on Status column (col index 9)
    for row_idx, t in enumerate(transports, start=1):
        bg = STATUS_BG.get(t.status)
        if bg:
            ts.append(("BACKGROUND", (9, row_idx), (9, row_idx), bg))

    table.setStyle(TableStyle(ts))
    elements.append(table)

    doc.build(elements)
    return response


@login_required
def send_password_otp(request):

    profile = UserProfile.objects.get(user=request.user)
    company = profile.company

    otp = random.randint(100000, 999999)

    request.session["password_change_otp"] = str(otp)

    subject = "Password Change OTP | CWS Inventory"

    html_message = render_to_string(
        "company/emails/reset_password_otp.html",
        {
            "name": company.name,
            "otp": otp,
        }
    )

    plain_message = strip_tags(html_message)

    email_msg = EmailMultiAlternatives(
        subject,
        plain_message,
        settings.EMAIL_HOST_USER,
        [company.email]
    )

    email_msg.attach_alternative(html_message, "text/html")

    logo_path = finders.find('accounts/images/logo.png')

    if logo_path:
        with open(logo_path, 'rb') as f:
            mime_image = MIMEImage(f.read())
            mime_image.add_header('Content-ID', '<logo>')
            mime_image.add_header('Content-Disposition', 'inline', filename='logo.png')
            email_msg.attach(mime_image)

    email_msg.send(fail_silently=False)

    return JsonResponse({
        "success": True
    })


import json

@login_required
def verify_password_otp(request):

    data = json.loads(request.body)

    otp = data.get("otp")

    session_otp = request.session.get(
        "password_change_otp"
    )

    if str(otp) == str(session_otp):

        request.session["otp_verified"] = True

        return JsonResponse({
            "success": True
        })

    return JsonResponse({
        "success": False,
        "message": "Invalid OTP"
    })
@login_required
def check_email(request):
    email = request.GET.get('email', '').strip()
    if not email:
        return JsonResponse({'exists': False})
    
    # Check if any user (other than current user) has this email
    from django.contrib.auth.models import User
    exists = User.objects.filter(email=email).exclude(id=request.user.id).exists()
    return JsonResponse({'exists': exists})

# --- OWNER VIEWS FOR STAFF INVENTORY OPERATIONS ---
from company.models import StaffActivityLog, StaffDeleteRequest
from inventory.models import Product, Category, PurchaseOrder, SalesOrder, Customer

def _is_owner(user):
    profile = UserProfile.objects.filter(user=user).first()
    return profile and profile.role == 'COMPANY_OWNER'

@login_required
def owner_staff_activity(request):
    if not _is_owner(request.user):
        messages.error(request, "Access denied.")
        return redirect('company_login')
        
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company
    
    logs = StaffActivityLog.objects.filter(company=company)
    
    # Optional filtering
    staff_id = request.GET.get('staff_id')
    if staff_id:
        logs = logs.filter(staff_id=staff_id)
        
    module = request.GET.get('module')
    if module:
        logs = logs.filter(module=module)
        
    action_type = request.GET.get('action_type')
    if action_type:
        logs = logs.filter(action_type=action_type)
        
    paginator = Paginator(logs, 50)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)
    
    staff_members = Staff.objects.filter(company=company)
    
    return render(request, 'company/staff_activity_log.html', {
        'company': company,
        'page_obj': page_obj,
        'staff_members': staff_members,
        'modules': StaffActivityLog.MODULE_CHOICES,
        'action_types': StaffActivityLog.ACTION_CHOICES,
    })

@login_required
def owner_delete_requests(request):
    if not _is_owner(request.user):
        messages.error(request, "Access denied.")
        return redirect('company_login')
        
    profile = UserProfile.objects.get(user=request.user)
    company = profile.company
    
    requests = StaffDeleteRequest.objects.filter(company=company, status='PENDING')
    
    return render(request, 'company/delete_requests.html', {
        'company': company,
        'requests': requests,
    })

@login_required
def owner_approve_delete(request, pk):
    if not _is_owner(request.user):
        messages.error(request, "Access denied.")
        return redirect('company_login')
        
    if request.method == 'POST':
        profile = UserProfile.objects.get(user=request.user)
        company = profile.company
        
        req = get_object_or_404(StaffDeleteRequest, pk=pk, company=company, status='PENDING')
        
        # Determine model class
        model_class = None
        if req.module == 'PRODUCT': model_class = Product
        elif req.module == 'CATEGORY': model_class = Category
        elif req.module == 'PO': model_class = PurchaseOrder
        elif req.module == 'SO': model_class = SalesOrder
        elif req.module == 'CUSTOMER': model_class = Customer
        
        if model_class:
            try:
                obj = model_class.objects.get(pk=req.object_id, company=company)
                obj.delete()
                req.status = 'APPROVED'
                req.reviewed_by = request.user
                req.reviewed_at = timezone.now()
                req.save()
                messages.success(request, f"Request approved. {req.module} '{req.object_repr}' deleted.")
            except model_class.DoesNotExist:
                messages.error(request, "Target object not found. It may have already been deleted.")
                req.status = 'REJECTED'
                req.rejection_note = "Object not found."
                req.reviewed_by = request.user
                req.reviewed_at = timezone.now()
                req.save()
        else:
            messages.error(request, "Unknown module.")
            
    return redirect('owner_delete_requests')

@login_required
def owner_reject_delete(request, pk):
    if not _is_owner(request.user):
        messages.error(request, "Access denied.")
        return redirect('company_login')
        
    if request.method == 'POST':
        profile = UserProfile.objects.get(user=request.user)
        company = profile.company
        
        req = get_object_or_404(StaffDeleteRequest, pk=pk, company=company, status='PENDING')
        
        note = request.POST.get('rejection_note', '')
        req.status = 'REJECTED'
        req.rejection_note = note
        req.reviewed_by = request.user
        req.reviewed_at = timezone.now()
        req.save()
        
        messages.success(request, f"Delete request for '{req.object_repr}' rejected.")
        
    return redirect('owner_delete_requests')
